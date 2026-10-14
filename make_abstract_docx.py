from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from pathlib import Path
p=Path(__file__).parent
text=(p/'abstract_en.txt').read_text(encoding='utf-8')
title=text.split('\n\nAbstract')[0]
body=text.split('Abstract\n\n')[1].split('\n\nKeywords:')[0]
keywords=text.split('Keywords: ')[1].strip()
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.67);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10)
style.paragraph_format.space_after=Pt(5)
p0=doc.add_paragraph();p0.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p0.add_run(title);r.bold=True;r.font.size=Pt(13)
p1=doc.add_paragraph();p1.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p1.add_run('Abstract');r.bold=True;r.font.size=Pt(11)
for para in body.split('\n\n'):
    p2=doc.add_paragraph(para);p2.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
p3=doc.add_paragraph();p3.add_run('Keywords: ').bold=True;p3.add_run(keywords)
doc.save(p/'abstract_en_one_page.docx')
